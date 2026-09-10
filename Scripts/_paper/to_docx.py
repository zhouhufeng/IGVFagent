#!/usr/bin/env python3
"""manuscript.html -> manuscript.docx

Walks the published manuscript and rebuilds it as a native Word document:
real heading styles, real tables, real numbered references, and inline
bold / italic / monospace / superscript preserved as runs rather than
flattened into plain text.
"""
import sys, re
from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK   = RGBColor(0x14,0x16,0x2A)
MUTED = RGBColor(0x5E,0x61,0x78)
INFER = RGBColor(0x3B,0x44,0xA6)
MEAS  = RGBColor(0xAF,0x5F,0x2B)
OK    = RGBColor(0x2C,0x6A,0x4C)
BODY, MONO, HEAD = "Calibri", "Consolas", "Calibri"


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd'); el.set(qn('w:val'),'clear'); el.set(qn('w:fill'),hexcolor)
    tcPr.append(el)


def add_runs(par, node, bold=False, italic=False, mono=False, sup=False, color=None):
    """Recursively emit runs, carrying inline formatting down the tree."""
    for child in node.children:
        if isinstance(child, NavigableString):
            txt = str(child)
            if not txt.strip() and not txt.startswith(" "):
                continue
            txt = re.sub(r"\s+", " ", txt)
            if not txt:
                continue
            r = par.add_run(txt)
            r.bold, r.italic = bold, italic
            r.font.name = MONO if mono else BODY
            r.font.size = Pt(9.5) if mono else Pt(10.5)
            if sup:
                r.font.superscript = True
                r.font.size = Pt(8)
            if color is not None:
                r.font.color.rgb = color
            continue
        n = child.name
        add_runs(par, child,
                 bold   = bold   or n in ("b","strong"),
                 italic = italic or n in ("em","i"),
                 mono   = mono   or n == "code" or "mono" in (child.get("class") or []),
                 sup    = sup    or n == "sup",
                 color  = MEAS if (child.get("class") or []) and "chg" in child.get("class")
                          else (OK if (child.get("class") or []) and "now" in child.get("class")
                          else (MUTED if (child.get("class") or []) and "was" in child.get("class") else color)))


def para(doc, node, style=None, space_after=6, size=10.5, color=None, italic=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    add_runs(p, node, italic=italic, color=color)
    for r in p.runs:
        if size != 10.5 and not r.font.superscript:
            r.font.size = Pt(size)
    return p


def build(src, out):
    soup = BeautifulSoup(open(src, encoding="utf-8").read(), "lxml")
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name, st.font.size = BODY, Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1.0)
        s.top_margin = s.bottom_margin = Inches(0.9)

    # ---- masthead ----
    t = soup.select_one("h1.title")
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(10)
    r = p.add_run(t.get_text(" ", strip=True))
    r.bold = True; r.font.size = Pt(19); r.font.name = HEAD; r.font.color.rgb = INK

    for sel, sz, col in (("p.byline",11,INK), ("p.corr",9.5,MUTED)):
        el = soup.select_one(sel)
        if el:
            pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(3)
            rr = pp.add_run(el.get_text(" ", strip=True))
            rr.font.size = Pt(sz); rr.font.color.rgb = col
            if sel == "p.corr": rr.font.name = MONO

    doc.add_paragraph()

    for sec in soup.select("section"):
        sid = sec.get("id","")
        h2 = sec.find("h2")
        if h2:
            num = h2.find("span", class_="num")
            label = num.get_text(strip=True) if num else ""
            if num: num.extract()
            title = h2.get_text(" ", strip=True)
            hp = doc.add_heading(level=1)
            hp.paragraph_format.space_before = Pt(16)
            hp.paragraph_format.space_after = Pt(6)
            hr = hp.add_run((f"{label}  " if label and label!="△" else "") + title)
            hr.font.name = HEAD; hr.font.size = Pt(15); hr.font.color.rgb = INK; hr.bold = True

        for el in sec.find_all(["h3","h4","p","ul","div","figure","ol"], recursive=True):
            cls = el.get("class") or []
            # skip nodes handled by an ancestor
            if el.find_parent(["figure","li"]) or (el.name=="p" and el.find_parent("div", class_=["abstract","note","card"])):
                if not (el.name=="p" and el.find_parent("div", class_="abstract")):
                    continue

            if el.name == "h3":
                hp = doc.add_heading(level=2); hp.paragraph_format.space_before = Pt(12)
                hp.paragraph_format.space_after = Pt(4)
                r = hp.add_run(el.get_text(" ", strip=True))
                r.font.name = HEAD; r.font.size = Pt(12); r.bold = True; r.font.color.rgb = INK
            elif el.name == "h4":
                hp = doc.add_paragraph(); hp.paragraph_format.space_before = Pt(11)
                hp.paragraph_format.space_after = Pt(3)
                r = hp.add_run(el.get_text(" ", strip=True).upper())
                r.font.name = MONO; r.font.size = Pt(8.5); r.bold = True; r.font.color.rgb = MEAS
            elif el.name == "p":
                if "eyebrow" in cls or "byline" in cls or "corr" in cls: continue
                para(doc, el, size=11.5 if "lede" in cls else 10.5,
                     color=MUTED if "sub" in cls else None,
                     italic="lede" in cls)
            elif el.name == "ul" and "edit-list" in cls:
                for li in el.find_all("li", recursive=False):
                    kind = li.find("span", class_="kind")
                    body = li.find("div")
                    pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(5)
                    pp.paragraph_format.left_indent = Inches(0.25)
                    if kind:
                        kr = pp.add_run(kind.get_text(strip=True).upper() + "  ")
                        kr.font.name = MONO; kr.font.size = Pt(8); kr.bold = True
                        kr.font.color.rgb = INFER if "fact" in (kind.get("class") or []) else MEAS
                    if body: add_runs(pp, body)
            elif el.name in ("ul","ol") and "refs" in cls:
                for i, li in enumerate(el.find_all("li", recursive=False), 1):
                    pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(3)
                    pp.paragraph_format.left_indent = Inches(0.3)
                    pp.paragraph_format.first_line_indent = Inches(-0.3)
                    nr = pp.add_run(f"{i}. "); nr.font.size = Pt(9.5)
                    add_runs(pp, li)
                    for r in pp.runs: r.font.size = Pt(9.5)
            elif el.name == "ul":
                for li in el.find_all("li", recursive=False):
                    pp = doc.add_paragraph(style="List Bullet")
                    pp.paragraph_format.space_after = Pt(3)
                    add_runs(pp, li)
            elif el.name == "div" and ("abstract" in cls or "note" in cls):
                lbl = el.find("span", class_="lbl")
                if lbl:
                    pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(2)
                    r = pp.add_run(lbl.get_text(strip=True).upper())
                    r.font.name = MONO; r.font.size = Pt(8); r.bold = True; r.font.color.rgb = INFER
                    lbl.extract()
                for sub in el.find_all("p", recursive=False):
                    pq = para(doc, sub, space_after=5)
                    pq.paragraph_format.left_indent = Inches(0.22)
            elif el.name == "figure":
                cap = el.find("figcaption")
                if cap:
                    pp = doc.add_paragraph(); pp.paragraph_format.space_before = Pt(12)
                    pp.paragraph_format.space_after = Pt(4)
                    add_runs(pp, cap)
                    for r in pp.runs: r.font.size = Pt(9)
                html_t = el.find("table")
                if html_t: emit_table(doc, html_t)
                for extra in el.find_all("figcaption")[1:]:
                    pp = doc.add_paragraph(); pp.paragraph_format.space_before = Pt(3)
                    add_runs(pp, extra)
                    for r in pp.runs: r.font.size = Pt(8.5); r.font.color.rgb = MUTED

    doc.save(out)
    return out


def emit_table(doc, html_t):
    heads = [th.get_text(" ", strip=True) for th in html_t.select("thead th")]
    rows = html_t.select("tbody tr")
    tbl = doc.add_table(rows=1, cols=max(1,len(heads)))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(heads):
        c = tbl.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h.upper())
        r.bold = True; r.font.size = Pt(7.5); r.font.name = MONO; r.font.color.rgb = MUTED
        shade(c, "F2F3F8")
    for tr in rows:
        cells = tr.find_all(["td","th"])
        row = tbl.add_row()
        is_grp = "grp" in (tr.get("class") or []) or "tot" in (tr.get("class") or [])
        for i, td in enumerate(cells[:len(tbl.columns)]):
            c = row.cells[i]; c.text = ""
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(1)
            add_runs(p, td)
            for r in p.runs:
                r.font.size = Pt(8.5)
                if is_grp: r.bold = True
            if is_grp: shade(c, "F2F3F8")
    for row in tbl.rows:
        for c in row.cells:
            c._tc.get_or_add_tcPr()


if __name__ == "__main__":
    print("wrote", build(sys.argv[1], sys.argv[2]))
