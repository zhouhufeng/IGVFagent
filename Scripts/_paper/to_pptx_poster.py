#!/usr/bin/env python3
"""poster.html -> poster.pptx  (one 48 x 36 in slide, editable)

PowerPoint's maximum slide dimension is 56 in, so 48 x 36 fits natively
and prints at true size with no scaling. Point sizes carry over from the
HTML unchanged, because both media measure type in points.
"""
import sys, re
from bs4 import BeautifulSoup, NavigableString
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

INK=RGBColor(0x14,0x16,0x2A); MUTED=RGBColor(0x5E,0x61,0x78)
INFER=RGBColor(0x3B,0x44,0xA6); MEAS=RGBColor(0xAF,0x5F,0x2B)
OK=RGBColor(0x2C,0x6A,0x4C); GROUND=RGBColor(0xF3,0xF4,0xF8)
SURF=RGBColor(0xFF,0xFF,0xFF); SURF2=RGBColor(0xF7,0xF8,0xFB); RULE=RGBColor(0xD6,0xD8,0xE4)
HDR_SUB=RGBColor(0xB9,0xBC,0xD8); HDR_DIM=RGBColor(0x9A,0x9E,0xC4); WHITE=RGBColor(0xFF,0xFF,0xFF)
SERIF="Georgia"; SANS="Calibri"; MONO="Consolas"
import os
SC=float(os.environ.get("POSTER_SCALE","1.0"))  # type scale

W,H = 48.0, 36.0
MX   = 0.85            # side margin (in)
HDRH = 5.4             # header band height
GAP  = 0.5
COLW = (W - 2*MX - 3*GAP)/4


def T(node):
    return re.sub(r"\s+"," ",node.get_text(" ",strip=True)).strip()


def tb(slide,x,y,w,h):
    t=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=t.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    return tf


def line(tf,first,text,size,color,font=SANS,bold=False,space=4,align=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after=Pt(space)
    r=p.add_run(); r.text=text
    r.font.size=Pt(size*SC); r.font.color.rgb=color; r.font.name=font; r.font.bold=bold
    if align: p.alignment=align
    return p


def runs(par,node,bold=False,italic=False,mono=False,size=15.5,color=INK):
    for ch in node.children:
        if isinstance(ch,NavigableString):
            t=re.sub(r"\s+"," ",str(ch))
            if not t.strip(): continue
            r=par.add_run(); r.text=t
            r.font.bold, r.font.italic = bold, italic
            r.font.name = MONO if mono else SANS
            r.font.size=Pt(size*SC); r.font.color.rgb=color
            continue
        cls=ch.get("class") or []
        runs(par,ch,bold or ch.name in ("b","strong"),
             italic or ch.name in ("em","i"),
             mono or ch.name=="code" or "mono" in cls,
             size, color)


def rect(slide,x,y,w,h,fill,line_col=None,rounded=False,accent=None):
    shp=slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x),Inches(y),Inches(w),Inches(h))
    if rounded: shp.adjustments[0]=0.03
    shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if line_col is None: shp.line.fill.background()
    else: shp.line.color.rgb=line_col; shp.line.width=Pt(1)
    shp.shadow.inherit=False
    if accent is not None:
        b=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(0.07),Inches(h))
        b.fill.solid(); b.fill.fore_color.rgb=accent; b.line.fill.background(); b.shadow.inherit=False
    return shp


def est_h(text,size,width_in,lead=1.42):
    """Rough wrapped height: chars per line from point size and column width."""
    size = size*SC
    cpl = max(12, int(width_in*96/(size*0.55)))
    lines = max(1, -(-len(text)//cpl))
    return lines*size*lead/72.0


def render_col(slide, col, x, y0, w):
    y=y0
    for el in col.find_all(True, recursive=False):
        cls=el.get("class") or []
        if el.name=="div" and "box" in cls:
            inner=[c for c in el.find_all(["span","h3","p","div","table"],recursive=False)]
            h=0.35
            for c in inner:
                if c.name=="span": h+=0.32
                elif c.name=="h3": h+=est_h(T(c),18,w-0.6)+0.1
                elif c.name=="p": h+=est_h(T(c),15.5,w-0.6)+0.12
                elif c.name=="div" and "row" in (c.get("class") or []): h+=1.55
            accent = INFER if "i" in cls else (MEAS if "m" in cls else None)
            rect(slide,x,y,w,h,SURF,RULE,rounded=True,accent=accent)
            yy=y+0.2
            for c in inner:
                if c.name=="span":
                    tf=tb(slide,x+0.3,yy,w-0.6,0.3)
                    line(tf,True,T(c).upper(),12,INFER if "i" in (c.get("class") or []) else MEAS,MONO,True,0)
                    yy+=0.32
                elif c.name=="h3":
                    hh=est_h(T(c),21,w-0.6)
                    tf=tb(slide,x+0.3,yy,w-0.6,hh)
                    line(tf,True,T(c),21,INK,SERIF,False,0); yy+=hh+0.1
                elif c.name=="p":
                    hh=est_h(T(c),15.5,w-0.6)
                    tf=tb(slide,x+0.3,yy,w-0.6,hh)
                    p=tf.paragraphs[0]; p.space_after=Pt(0); runs(p,c,size=15.5)
                    yy+=hh+0.12
                elif c.name=="div" and "row" in (c.get("class") or []):
                    yy=render_row(slide,c,x+0.3,yy,w-0.6)+0.1
            y+=h+0.34
        elif el.name=="div":
            y=render_block(slide,el,x,y,w)
        else:
            y=render_node(slide,el,x,y,w)
    return y


def render_row(slide,row,x,y,w):
    kids=row.find_all("div",recursive=False)
    n=max(1,len(kids)); cw=(w-0.3*(n-1))/n; maxb=y
    for i,d in enumerate(kids):
        cx=x+i*(cw+0.3)
        sv=d.find("div",class_="stat"); sl=d.find("div",class_="statlbl")
        yy=cy=y
        if "box" in (d.get("class") or []):
            rect(slide,cx,y,cw,1.85,SURF,RULE,rounded=True); yy=y+0.18; cx+=0.22; cw2=cw-0.44
        else: cw2=cw
        lb=d.find("span",class_="lbl")
        if lb:
            tf=tb(slide,cx,yy,cw2,0.3)
            line(tf,True,T(lb).upper(),12,INFER if "i" in (lb.get("class") or []) else MEAS,MONO,True,0)
            yy+=0.32
        if sv:
            sz=40 if "sm" in (sv.get("class") or []) else 56
            tf=tb(slide,cx,yy,cw2,sz/52.0)
            line(tf,True,T(sv),sz,MEAS if "m" in (sv.get("class") or []) else INFER,SERIF,False,0)
            yy+=sz/58.0+0.12
        if sl:
            hh=est_h(T(sl),12,cw2)
            tf=tb(slide,cx,yy,cw2,hh)
            line(tf,True,T(sl).upper(),12,MUTED,MONO,False,0); yy+=hh
        maxb=max(maxb,yy)
    return maxb+0.1


def render_block(slide,el,x,y,w):
    cls=el.get("class") or []
    if "row" in cls: return render_row(slide,el,x,y,w)
    for kid in el.find_all(True,recursive=False):
        y=render_node(slide,kid,x,y,w)
    return y


def render_node(slide,el,x,y,w):
    cls=el.get("class") or []
    if el.name=="h2":
        n=el.find("span",class_="n")
        if n:
            tf=tb(slide,x,y,w,0.3); line(tf,True,T(n).upper(),14,MEAS,MONO,True,0); y+=0.34
            n.extract()
        t=T(el); hh=est_h(t,31,w,1.15)
        tf=tb(slide,x,y,w,hh); line(tf,True,t,31,INK,SERIF,False,0)
        y+=hh+0.1
        ln=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(0.035))
        ln.fill.solid(); ln.fill.fore_color.rgb=INK; ln.line.fill.background(); ln.shadow.inherit=False
        return y+0.26
    if el.name=="h3":
        hh=est_h(T(el),18,w); tf=tb(slide,x,y,w,hh)
        line(tf,True,T(el),18,INK,SANS,True,0); return y+hh+0.12
    if el.name=="p":
        t=T(el)
        if not t: return y
        size=13 if "muted" in cls else 15.5
        hh=est_h(t,size,w); tf=tb(slide,x,y,w,hh)
        p=tf.paragraphs[0]; p.space_after=Pt(0)
        runs(p,el,size=size,color=MUTED if "muted" in cls else INK)
        return y+hh+0.15
    if el.name=="ul":
        items=el.find_all("li",recursive=False); tot=0
        for li in items: tot+=est_h(T(li),15,w-0.25)+0.1
        tf=tb(slide,x,y,w,tot)
        for i,li in enumerate(items):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.space_after=Pt(7)
            b=p.add_run(); b.text="— "
            b.font.color.rgb=INFER if "i" in cls else MEAS; b.font.size=Pt(15*SC); b.font.name=SANS
            runs(p,li,size=15)
        return y+tot+0.15
    if el.name=="table":
        heads=[T(th) for th in el.select("thead th")]; body=el.select("tbody tr")
        rows=len(body)+1; cols=max(1,len(heads)); rh=0.36*SC
        tbl=slide.shapes.add_table(rows,cols,Inches(x),Inches(y),Inches(w),Inches(rh*rows)).table
        for i,h in enumerate(heads):
            c=tbl.cell(0,i); c.text=""
            r=c.text_frame.paragraphs[0].add_run(); r.text=h.upper()
            r.font.size=Pt(11.5*SC); r.font.bold=True; r.font.name=MONO; r.font.color.rgb=MUTED
            c.fill.solid(); c.fill.fore_color.rgb=SURF2
        for ri,tr in enumerate(body,1):
            tds=tr.find_all(["td","th"]); istot="tot" in (tr.get("class") or [])
            for ci,td in enumerate(tds[:cols]):
                c=tbl.cell(ri,ci); c.text=""
                p=c.text_frame.paragraphs[0]; runs(p,td,size=14)
                for r in p.runs:
                    r.font.size=Pt(14*SC)
                    if istot: r.font.bold=True
                    if "yes" in (td.get("class") or []): r.font.color.rgb=OK; r.font.bold=True
                c.fill.solid(); c.fill.fore_color.rgb=SURF2 if istot else SURF
        for r in tbl.rows: r.height=Inches(rh)
        return y+rh*rows+0.24
    if el.name=="div" and "term" in cls:
        lines=[l for l in el.get_text("\n").split("\n") if l.strip()]
        h=0.245*SC*len(lines)+0.4
        rect(slide,x,y,w,h,INK,None,rounded=True)
        tf=tb(slide,x+0.28,y+0.2,w-0.56,h-0.4)
        for i,l in enumerate(lines):
            col=(RGBColor(0x8C,0x95,0xEC) if l.strip().startswith(">") else
                 RGBColor(0x63,0xB9,0x8C) if l.strip().startswith("$") else
                 RGBColor(0x9A,0x9C,0xB4) if l.strip().startswith("#") else
                 RGBColor(0xDE,0x9A,0x62))
            line(tf,i==0,l,12.5,col,MONO,False,1)
        return y+h+0.22
    if el.name=="div" and "flow" in cls:
        sts=el.find_all("div",class_="st"); n=max(1,len(sts))
        cw=(w-0.28*(n-1))/n
        for i,st in enumerate(sts):
            cx=x+i*(cw+0.28)
            rect(slide,cx,y,cw,1.5,SURF,RULE,rounded=True)
            k=st.find("span",class_="k"); v=st.find("span",class_="v")
            tf=tb(slide,cx+0.14,y+0.18,cw-0.28,1.2)
            if k: line(tf,True,T(k).upper(),11,INFER,MONO,True,4)
            if v: line(tf,not k,T(v),14,INK,SANS,True,0,PP_ALIGN.CENTER)
        return y+1.72
    if el.name=="div":
        return render_block(slide,el,x,y,w)
    return y


def main(src,out):
    soup=BeautifulSoup(open(src,encoding="utf-8").read(),"lxml")
    prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H)
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb=GROUND

    # header band
    rect(s,0,0,W,HDRH,INK,None)
    tf=tb(s,MX,0.62,W*0.62,3.6)
    line(tf,True,"IGVF Agent",76/SC,WHITE,SERIF,False,10)
    line(tf,False,T(soup.select_one("p.sub")),30/SC,HDR_SUB,SERIF,False,14)
    tf2=tb(s,MX,4.15,W*0.62,1.0)
    line(tf2,True,T(soup.select_one("p.authors")),17/SC,RGBColor(0xD6,0xD8,0xEC),SANS,False,5)
    line(tf2,False,T(soup.select_one("p.affil")),14/SC,HDR_DIM,SANS,False,0)
    tf3=tb(s,W-MX-12.5,1.5,12.5,2.6)
    line(tf3,True,"github.com/zhouhufeng/IGVFagent",20/SC,WHITE,MONO,False,7,PP_ALIGN.RIGHT)
    line(tf3,False,"Apache-2.0 · pip installable · containerized",15/SC,HDR_SUB,MONO,False,10,PP_ALIGN.RIGHT)
    line(tf3,False,"60 SKILLS · 175 TOOLS · 6 ARCHIVES",13/SC,RGBColor(0xC9,0xCC,0xE6),MONO,True,0,PP_ALIGN.RIGHT)

    cols=soup.select("main > div.col")
    top=HDRH+0.55
    bottoms=[]
    for i,c in enumerate(cols):
        x=MX+i*(COLW+GAP)
        bottoms.append(render_col(s,c,x,top,COLW))

    # footer
    fy=H-0.75
    ln=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(MX),Inches(fy),Inches(W-2*MX),Inches(0.025))
    ln.fill.solid(); ln.fill.fore_color.rgb=RULE; ln.line.fill.background(); ln.shadow.inherit=False
    tf=tb(s,MX,fy+0.16,W/2,0.35)
    line(tf,True,"IGVF Agent · Zhou, Ma, Yang, … , Lin · Harvard T.H. Chan School of Public Health",13/SC,MUTED,MONO,False,0)
    tf=tb(s,W/2,fy+0.16,W/2-MX,0.35)
    line(tf,True,"github.com/zhouhufeng/IGVFagent · Apache-2.0",13/SC,MUTED,MONO,False,0,PP_ALIGN.RIGHT)

    prs.save(out)
    print(f"wrote {out} — 1 slide, {W:.0f} x {H:.0f} in")
    print("column bottoms (in):", [round(b,2) for b in bottoms], "| usable to", H-0.9)


if __name__=="__main__":
    main(sys.argv[1],sys.argv[2])
